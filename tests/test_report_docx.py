"""Tests for phase0_foundations/report_docx.py (docx working paper + email body)."""

from __future__ import annotations

from docx import Document

from phase0_foundations.models import ExceptionItem, VerificationRun
from phase0_foundations.report_docx import build_report_docx, email_body_html


def _run() -> VerificationRun:
    return VerificationRun(
        id="abc123",
        status="done",
        inputs={
            "borrower": "leasy",
            "coverage_pct": 100.0,
            "independent_source_present": True,
            "bank_statement_count": 207,
            "statement_breakdown": {
                "totals": {"bank_rows": 207, "bank_in": 8119984.02,
                           "bank_out": 7075788.89, "bank_cash_total": 15195772.91,
                           "fx_base_currency": "USD"},
                "by_category": {"kushki (payment processor)": {"rows": 148, "in": 8105501.22, "out": 0.0}},
                "by_month": {},
                "by_statement": {},
            },
        },
        aggregates={"collections": 80767532.26, "calculated_collections": 8119984.02,
                    "fx_base_currency": "USD"},
        exceptions=[
            ExceptionItem(id="abc123:recon:collections", kind="reconciliation",
                          severity=1.0, status="pending",
                          description="Collections variance 894.7%")],
    )


def test_build_report_docx_writes_valid_word_document(tmp_path):
    run = _run()
    path = build_report_docx(tmp_path, run)
    assert path.name == "run_abc123.docx"
    assert path.exists()

    doc = Document(str(path))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Verification Report" in text
    assert "leasy" in text
    assert "abc123" in text
    # tables render (metadata + statement breakdown present)
    assert len(doc.tables) >= 2


def test_email_body_html_summarises_run():
    html = email_body_html(_run())
    assert "<html" in html
    assert "leasy" in html
    assert "abc123" in html
    assert "8,119,984.02" in html          # calculated collections
    assert "80,767,532.26" in html         # reported collections
    assert "15,195,772.91" in html         # statement total
    assert "1" in html                     # exceptions count
